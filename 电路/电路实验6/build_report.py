from copy import deepcopy
from pathlib import Path

import pypdfium2 as pdfium
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).parent
TEMPLATE = ROOT / "电路技术基础（上）实验报告模板.docx"
OUT = ROOT / "实验六 基尔霍夫定律验证实验报告.docx"


def make_figures():
    """Render and crop the three original circuit schematics from page 47."""
    page = pdfium.PdfDocument(str(ROOT / "电路原理实验指导书_46-48.pdf"))[1]
    image = page.render(scale=3).to_pil()
    # The left half of page 47 contains the original circuit diagrams (not redrawn).
    crops = {
        "图1_指导书电路.png": (230, 155, 765, 560),
        "图2_指导书电路.png": (225, 540, 770, 995),
        "图3_指导书电路.png": (215, 950, 785, 1415),
    }
    for name, box in crops.items():
        image.crop(box).save(ROOT / name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    item = OxmlElement("w:tblHeader")
    item.set(qn("w:val"), "true")
    tr_pr.append(item)


def set_table_grid(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement("w:" + edge)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), "808080")
        borders.append(element)
    tbl_pr.append(borders)


def style_run(run, size=10.5, bold=False):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def add_text(doc, text="", *, size=10.5, bold=False, align=None, first_indent=True, space_after=3):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(space_after)
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    style_run(r, size, bold)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    style_run(r, 14 if level == 1 else 12, True)
    return p


def add_equation(doc, text):
    return add_text(doc, text, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, space_after=4)


def add_figure(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(12.8))
    add_text(doc, caption, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, space_after=6)


def add_placeholder(doc, title):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_grid(table)
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.width = Cm(15.5)
    set_cell_shading(cell, "F2F2F2")
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(26)
    r = p.add_run("【待插入：" + title + "】\n请粘贴对应的手工原始数据记录照片，保持读数、单位和测量方向清晰可辨。")
    style_run(r, 10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_results_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_grid(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "D9EAF7")
        for run in cell.paragraphs[0].runs:
            style_run(run, 9.5, True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for run in p.runs:
                    style_run(run, 9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def delete_after(paragraph):
    element = paragraph._element
    parent = element.getparent()
    seen = False
    for child in list(parent):
        if seen:
            # Keep the section properties; python-docx needs them to add tables.
            if child.tag != qn("w:sectPr"):
                parent.remove(child)
        if child is element:
            seen = True


def build():
    make_figures()
    doc = Document(str(TEMPLATE))
    # Retain the template cover; replace only the unfilled report body.
    cover = doc.paragraphs
    cover[2].text = "实验报告"
    for r in cover[2].runs:
        style_run(r, 22, True)
    cover[10].text = "实验名称：基尔霍夫定律的验证"
    cover[11].text = "实验日期：________________"
    cover[12].text = "指导教师：黄毅刚老师"
    for i in (10, 11, 12):
        for r in cover[i].runs:
            style_run(r, 12)
    delete_after(cover[14])
    doc.add_page_break()

    add_heading(doc, "一、实验目的")
    add_text(doc, "1. 加深对基尔霍夫电流定律（KCL）和基尔霍夫电压定律（KVL）的理解，掌握其在直流集中参数电路中的验证方法。")
    add_text(doc, "2. 学习根据电路结构选取节点、支路和独立回路，规定电流与电压的参考方向，并利用实测量进行代数检验。")
    add_text(doc, "3. 熟悉直流稳压电源、直流电流表和直流电压表的使用，认识仪表内阻及读数精度对验证结果的影响。")

    add_heading(doc, "二、实验仪器与设备")
    add_results_table(doc, ["仪器/器材", "主要规格或用途", "数量"], [
        ["电路原理实验箱", "含验证基尔霍夫定律专用线路与开关", "1 台"],
        ["直流稳压电源", "0～24 V、+12 V、5 V 输出", "1 台"],
        ["直流电流表", "0～30 mA，用于支路电流测量", "3 只"],
        ["直流电压表", "0～30 V，用于元件端电压测量", "2 只"],
        ["连接导线", "实验箱专用导线", "若干"],
    ])

    add_heading(doc, "三、实验原理")
    add_text(doc, "对于任一集总参数电路的节点，流入节点的电流代数和等于流出节点的电流代数和，即 Σi＝0；等价地，ΣI入＝ΣI出。这就是基尔霍夫电流定律（KCL）。")
    add_text(doc, "对于任一闭合回路，沿预先选定的绕行方向，各支路电压的代数和为零，即 Σu＝0。穿越元件时，电压参考方向与绕行方向相同取正、相反取负（也可采用等价但须全程一致的符号约定）。这就是基尔霍夫电压定律（KVL）。")
    add_text(doc, "本实验分别采用指导书图1、图2、图3所示电路。图1为单回路双电源电路，适合验证串联支路电流关系及回路电压关系；图2为一节点分流、两支路并联结构；图3为双电源、三支路汇合结构。以下计算均以原始记录的方向标注为准。")

    add_heading(doc, "四、实验电路、节点与参考方向")
    add_figure(doc, ROOT / "图1_指导书电路.png", "图1  指导书中的验证电路一（原图）")
    add_text(doc, "图1取全电路为验证回路。主支路电流依次在 a—R1、R1—R2、R2—R5 处测量；按图中箭头绕行，两个电源与 R1、R2、R5 的端电压参与KVL检验。")
    add_figure(doc, ROOT / "图2_指导书电路.png", "图2  指导书中的验证电路二（原图）")
    add_text(doc, "图2取 R1、R2、R3 汇合处为考察节点，约定 A1 流入节点、A2 与 A3 流出节点；左回路含 Usa、R1、R3，右回路由 R2、R3 构成。")
    add_figure(doc, ROOT / "图3_指导书电路.png", "图3  指导书中的验证电路三（原图）")
    add_text(doc, "图3取 R3、R6、R2 汇合处为考察节点，约定 A1、A2 由两侧流入，A3 经 R6 向下流出。左回路包含 Usa、R1、R3、R6；右回路包含 Usb、R5、R2、R6。")

    add_heading(doc, "五、实验步骤")
    add_text(doc, "1. 断电接线，按指导书选定图1、图2、图3电路；检查电源极性、开关位置和电流表串联/电压表并联方式。")
    add_text(doc, "2. 在接通电源前标出各支路电流和元件电压的参考方向，选择合适量程；通电后先观察读数稳定性，再记录读数及其单位。")
    add_text(doc, "3. 对每个所选节点测量相关支路电流，对每个所选回路测量电源电压和元件电压；实验结束后断开电源、整理线路。")

    add_heading(doc, "六、原始记录与基尔霍夫定律验证")
    add_text(doc, "说明：按要求，原始数据表不在本报告中重新录入。请将手工记录数据照片粘贴至下列占位框。后续验算数值依据当前提供的电子转录稿，最终以粘贴的原始照片为准。")

    add_heading(doc, "6.1 图1：串联支路与单回路", level=2)
    add_placeholder(doc, "图1手工测量数据表照片")
    add_text(doc, "KCL检验：串联电路任一截面电流应相等。电子转录值为 15.698 mA、15.645 mA 和 15.000 mA；前两处之差为 0.053 mA（约0.34%），第三处与第一处最大相差 0.698 mA（约4.45%）。前两项符合较好，第三项应结合原始照片复核小数位、量程及读数。")
    add_equation(doc, "I(a—R1) ≈ I(R1—R2) ≈ I(R2—R5)")
    add_text(doc, "KVL检验：取与图示一致的绕行方向，Usa 的电压上升抵消 Usb 与三个电阻的电压降。")
    add_equation(doc, "ΔU₁＝Usa－Usb－U₁－U₂－U₅＝10.071－5.618－1.529－1.532－1.390＝0.002 V")
    add_text(doc, "残差仅为 2 mV，相对 Usa 约为0.020%，在读数分辨率范围内，图1的KVL得到很好验证。")

    add_heading(doc, "6.2 图2：分流节点与双回路", level=2)
    add_placeholder(doc, "图2手工测量数据表照片")
    add_text(doc, "KCL检验：以流入节点为正，ΔI₂＝A1－A2－A3＝30.3－15.4－15.3＝－0.4 mA。相对流入电流的闭合误差为 |ΔI₂|/A1×100%＝1.32%，可认为节点电流守恒得到验证。")
    add_equation(doc, "A1 ≈ A2＋A3：30.3 mA ≈ 15.4 mA＋15.3 mA")
    add_text(doc, "KVL检验：左回路采用 Usa、R1、R3；右回路采用并联支路 R2、R3。")
    add_results_table(doc, ["回路", "代数残差", "相对误差", "结论"], [
        ["左回路 I", "10.080－6.681－3.354＝0.045 V", "0.45%（相对 Usa）", "满足KVL"],
        ["右回路 II", "3.356－3.354＝0.002 V", "0.06%（相对约3.35 V）", "满足KVL"],
    ])

    add_heading(doc, "6.3 图3：双电源三支路节点与回路", level=2)
    add_placeholder(doc, "图3手工测量数据表照片")
    add_text(doc, "KCL检验：按 A1、A2 流入而 A3 流出的约定，ΔI₃＝A1＋A2－A3＝13.230＋2.496－15.840＝－0.114 mA；相对 A3 的闭合误差为0.72%，节点电流关系成立。")
    add_equation(doc, "A1＋A2 ≈ A3：13.230 mA＋2.496 mA ≈ 15.840 mA")
    add_text(doc, "KVL检验：以电子文本中“R6＝3.833 V、R5＝0.495 V、Usa＝10.080 V、Usb＝5.040 V”的一组转录为例：")
    add_results_table(doc, ["回路", "代数残差", "相对误差", "结论"], [
        ["左回路 I", "10.080－2.960－2.971－3.833＝0.316 V", "3.13%（相对 Usa）", "近似满足KVL"],
        ["右回路 II", "5.040－0.495－0.546－3.833＝0.166 V", "3.29%（相对 Usb）", "近似满足KVL"],
    ])
    add_text(doc, "数据复核说明：另一份电子转录将 R6 记为 3.733 V、R5 记为 0.504 V。若手工照片确认该组读数，左、右回路残差应分别改为 0.416 V 和 0.257 V。该差异说明图3必须以粘贴的手工记录为准后定稿，不能将两份转录混用。")

    add_heading(doc, "七、误差分析")
    add_text(doc, "1. 仪表内阻引入系统误差。电流表串入支路会增加支路电阻，改变原电路工作点；提供的电表数据表明，20 mA 挡内阻约为10.41 Ω，已约为220 Ω电阻的4.7%，对图2、图3的支路分流尤其不能忽略。电压表在2 V及以上量程内阻约为2.990 MΩ，远大于220 Ω元件，负载效应相对很小。")
    add_text(doc, "2. 元件与电源的非理想性。电阻标称值有允差，导线、接触点和开关有接触电阻；稳压电源带载后会出现内阻压降和显示读数波动。这些因素均会使KVL残差不完全为零。")
    add_text(doc, "3. 量程、分辨率与人为读数。电流和电压显示的末位量化误差、换挡后未充分稳定、正负端或参考方向标注不一致，都会在多个读数相减时累积放大。图1第三个电流记录的有效数字与前两项不一致，是应首先核查的项目。")
    add_text(doc, "4. 图3的数据一致性。两份电子转录对 R6（3.733 V/3.833 V）和 R5（0.504 V/0.495 V）不一致，使KVL残差发生明显改变。这属于记录或转录误差，不应被解释为定律失效；应以带单位、挡位和极性的手工原始记录照片复核。")

    add_heading(doc, "八、实验结论与思考")
    add_text(doc, "1. 图1的单回路电压代数和残差为0.002 V，KVL验证效果良好；其串联电流前两处差异仅0.34%，第三处读数需复核。")
    add_text(doc, "2. 图2节点电流闭合误差为1.32%，两个回路电压残差分别为0.045 V和0.002 V，均支持KCL、KVL在该直流电路中成立。")
    add_text(doc, "3. 图3节点电流闭合误差为0.72%，KVL残差约为3%量级；考虑电流表内阻、元件允差及当前转录不一致，结果仍体现定律规律，但必须在粘贴原始记录后复算并统一最终数值。")
    add_text(doc, "思考题：基尔霍夫定律适用于任何集总参数电路，与元件是否线性、是否含源无关。若电流表内阻相对支路电阻较大，会显著改变支路电流；若电压表内阻过小，会产生分流并改变被测电压，因此选择合适量程和高输入阻抗电压表是提高验证精度的关键。")

    # Page setup suitable for the report body.
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    build()
